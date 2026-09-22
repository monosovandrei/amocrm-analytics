#!/usr/bin/env python3
"""Private, offline document extraction. One bounded JSON request on stdin.

Requires Python 3.10+, Poppler (PDF), Tesseract rus+eng (scanned PDF) and
openpyxl (XLSX). DOCX uses only the standard library. No packages are downloaded.
Run under a dedicated, network-disabled service; POSIX resource limits are
also applied by the CLI. stdout contains metadata only, never document text.
"""

from __future__ import annotations

import csv
import contextlib
import hashlib
import io
import json
import math
import os
import re
import shutil
import signal
import stat
import subprocess
import sys
import tempfile
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

EXTRACTOR_VERSION = "local-documents-v1"
MAX_REQUEST = 16 * 1024
MAX_FILE = 20 * 1024 * 1024
MAX_ZIP_EXPANDED = 64 * 1024 * 1024
MAX_XML = 16 * 1024 * 1024
MAX_OUTPUT = 16 * 1024 * 1024
MAX_PAGES = 100
MAX_TEXT = 200_000
MAX_UNITS = 50_000
MAX_CELLS = 50_000
TIMEOUT_SECONDS = 120
MEMORY_BYTES = 768 * 1024 * 1024
MIMES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class ExtractionError(Exception):
    """Only fixed, non-sensitive error codes cross the CLI boundary."""


class Budget:
    def __init__(self, seconds: float = TIMEOUT_SECONDS):
        self.deadline = time.monotonic() + seconds

    def remaining(self) -> float:
        remaining = self.deadline - time.monotonic()
        if remaining <= 0:
            raise ExtractionError("TIME_LIMIT")
        return remaining


class Result:
    def __init__(self, sha256: str, kind: str):
        self.sha256 = sha256
        self.kind = kind
        self.units: list[dict[str, Any]] = []
        self.problems: list[str] = []
        self.text_chars = 0

    def problem(self, code: str):
        if code not in self.problems:
            self.problems.append(code)

    def add(self, text: str, locator: dict, **extra) -> bool:
        if not text:
            return True
        if len(self.units) >= MAX_UNITS:
            self.problem("UNIT_LIMIT")
            return False
        remaining = MAX_TEXT - self.text_chars
        if remaining <= 0:
            self.problem("TEXT_LIMIT")
            return False
        clipped = len(text) > remaining
        if clipped:
            text = text[:remaining]
            self.problem("TEXT_LIMIT")
        self.units.append({"locator": locator, "text": text,
                           "complete": not clipped, **extra})
        self.text_chars += len(text)
        return not clipped

    def payload(self) -> dict:
        if not self.units:
            self.problem("NO_TEXT")
        return {
            "extractorVersion": EXTRACTOR_VERSION,
            "sourceSha256": self.sha256,
            "format": self.kind,
            "mimeType": MIMES.get(self.kind),
            # COMPLETE concerns extraction only; it never asserts sent-history coverage.
            "status": "UNVERIFIED" if self.problems else "COMPLETE",
            "problems": self.problems,
            "textChars": self.text_chars,
            "units": self.units,
        }


def bounded_read(stream, limit: int) -> bytes:
    data = stream.read(limit + 1)
    if len(data) > limit:
        raise ExtractionError("SIZE_LIMIT")
    return data


def canonical_json(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, allow_nan=False,
                      separators=(",", ":"), sort_keys=True).encode("utf-8")


def private_directory(path: Path):
    if path.is_symlink():
        raise ExtractionError("UNSAFE_PATH")
    path.mkdir(mode=0o700, exist_ok=True)
    if not path.is_dir():
        raise ExtractionError("UNSAFE_PATH")
    if os.name == "posix":
        os.chmod(path, 0o700)


def checked_source(request: dict) -> tuple[Path, Path, str, str | None]:
    if not isinstance(request, dict) or set(request) != {"filePath", "sha256", "mimeType"}:
        raise ExtractionError("INVALID_REQUEST")
    filename, digest, mime = request["filePath"], request["sha256"], request["mimeType"]
    if (not isinstance(filename, str) or len(filename) > 4096 or "\x00" in filename
            or not isinstance(digest, str) or not re.fullmatch("[a-fA-F0-9]{64}", digest)
            or (mime is not None and (not isinstance(mime, str) or len(mime) > 256))):
        raise ExtractionError("INVALID_REQUEST")
    configured = os.environ.get("CRM_CONTROL_DOCUMENT_DIR", "")
    root, candidate = Path(configured), Path(filename)
    # UNC / double-slash roots can access a remote share on Windows.
    if (not configured or not root.is_absolute() or not candidate.is_absolute()
            or configured.startswith(("//", "\\\\")) or filename.startswith(("//", "\\\\"))
            or ".." in candidate.parts or root.is_symlink()):
        raise ExtractionError("UNSAFE_PATH")
    root = root.resolve(strict=True)
    try:
        relative = candidate.relative_to(root)
    except ValueError:
        raise ExtractionError("OUTSIDE_ARCHIVE") from None
    current = root
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            raise ExtractionError("UNSAFE_PATH")
    if not relative.parts or not candidate.is_file():
        raise ExtractionError("INVALID_SOURCE")
    return root, candidate, digest.lower(), mime.split(";", 1)[0].strip().lower() if mime else None


def copy_verified(source: Path, target: Path, digest: str, budget: Budget):
    descriptor = os.open(source, os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
                         | getattr(os, "O_BINARY", 0))
    with os.fdopen(descriptor, "rb") as incoming, target.open("xb") as outgoing:
        metadata = os.fstat(incoming.fileno())
        if not stat.S_ISREG(metadata.st_mode) or metadata.st_size > MAX_FILE:
            raise ExtractionError("FILE_LIMIT")
        hasher, count = hashlib.sha256(), 0
        while True:
            budget.remaining()
            block = incoming.read(64 * 1024)
            if not block:
                break
            count += len(block)
            if count > MAX_FILE:
                raise ExtractionError("FILE_LIMIT")
            hasher.update(block)
            outgoing.write(block)
        if hasher.hexdigest() != digest:
            raise ExtractionError("HASH_MISMATCH")
    if os.name == "posix":
        os.chmod(target, 0o600)


def parse_xml(data: bytes, *, generated=False):
    # ElementTree does not fetch external entities. Reject document DTDs outright.
    upper = data.upper()
    if b"<!ENTITY" in upper or (not generated and b"<!DOCTYPE" in upper):
        raise ExtractionError("UNSAFE_XML")
    return ET.fromstring(data)


def inspect_zip(path: Path) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(path)
    try:
        entries = archive.infolist()
        if len(entries) > 4000 or len({item.filename for item in entries}) != len(entries):
            raise ExtractionError("ZIP_LIMIT")
        expanded = 0
        for item in entries:
            expanded += item.file_size
            if (expanded > MAX_ZIP_EXPANDED or item.file_size > MAX_XML
                    or item.flag_bits & 1
                    or (item.file_size > 1024 * 1024
                        and item.file_size > max(item.compress_size, 1) * 1000)):
                raise ExtractionError("ZIP_LIMIT")
            if (item.filename.startswith(("/", "\\")) or "\\" in item.filename
                    or ".." in item.filename.split("/")):
                raise ExtractionError("UNSAFE_ZIP")
        return archive
    except BaseException:
        archive.close()
        raise


def identify(path: Path) -> str:
    with path.open("rb") as stream:
        magic = stream.read(8)
    if magic.startswith(b"%PDF-"):
        return "pdf"
    if magic.startswith(b"PK"):
        with inspect_zip(path) as archive:
            names = set(archive.namelist())
            word, excel = "word/document.xml" in names, "xl/workbook.xml" in names
            if word != excel:
                return "docx" if word else "xlsx"
    return "unsupported"


def run_tool(name: str, arguments: list[str], work: Path, budget: Budget) -> bytes:
    executable = shutil.which(name)
    if not executable:
        raise ExtractionError("DEPENDENCY_" + name.upper() + "_MISSING")
    environment = {key: os.environ[key] for key in ("PATH", "SYSTEMROOT", "WINDIR")
                   if key in os.environ}
    environment.update({"LC_ALL": "C", "LANG": "C", "OMP_THREAD_LIMIT": "1",
                        "HOME": str(work), "TMPDIR": str(work), "TEMP": str(work)})
    with tempfile.TemporaryFile(dir=work) as output:
        process = subprocess.Popen([executable, *arguments], cwd=work, env=environment,
                                   stdin=subprocess.DEVNULL, stdout=output,
                                   stderr=subprocess.DEVNULL,
                                   start_new_session=os.name == "posix")
        try:
            while process.poll() is None:
                if os.fstat(output.fileno()).st_size > MAX_OUTPUT:
                    raise ExtractionError("TOOL_OUTPUT_LIMIT")
                try:
                    process.wait(timeout=min(0.1, budget.remaining()))
                except subprocess.TimeoutExpired:
                    pass
        except BaseException:
            stop_process(process)
            raise
        if process.returncode:
            raise ExtractionError("TOOL_" + name.upper() + "_FAILED")
        output.seek(0)
        return bounded_read(output, MAX_OUTPUT)


def stop_process(process: subprocess.Popen):
    if process.poll() is None:
        try:
            if os.name == "posix":
                os.killpg(process.pid, signal.SIGKILL)
            else:
                process.kill()
        except ProcessLookupError:
            pass
    process.wait()


def rectangle(element) -> list[float] | None:
    values = [float(element.attrib[key]) for key in ("xMin", "yMin", "xMax", "yMax")]
    if not all(math.isfinite(value) for value in values):
        raise ExtractionError("INVALID_COORDINATES")
    if values[2] < values[0] or values[3] < values[1]:
        raise ExtractionError("INVALID_COORDINATES")
    return values


def ocr_page(source: Path, page: int, work: Path, budget: Budget, result: Result):
    image = work / "page"
    run_tool("pdftoppm", ["-f", str(page), "-l", str(page), "-singlefile",
                        "-scale-to", "2500", "-r", "200", "-gray", "-png",
                        str(source), str(image)], work, budget)
    image_path = image.with_suffix(".png")
    if not image_path.is_file() or image_path.stat().st_size > MAX_OUTPUT:
        raise ExtractionError("RASTER_LIMIT")
    data = run_tool("tesseract", [str(image_path), "stdout", "-l", "rus+eng",
                                "--oem", "1", "tsv"], work, budget)
    result.problem("OCR_UNVERIFIED")
    groups: dict[tuple, list[dict]] = {}
    for row in csv.DictReader(io.StringIO(data.decode("utf-8")), delimiter="\t"):
        budget.remaining()
        if row.get("level") != "5" or not row.get("text", "").strip():
            continue
        key = tuple(row[field] for field in ("block_num", "par_num", "line_num"))
        x, y, w, h = (int(row[field]) for field in ("left", "top", "width", "height"))
        confidence = float(row["conf"])
        if min(x, y, w, h) < 0 or not math.isfinite(confidence):
            raise ExtractionError("INVALID_OCR_OUTPUT")
        groups.setdefault(key, []).append({"text": row["text"], "x": x, "y": y,
                                          "w": w, "h": h, "confidence": confidence})
    for index, words in enumerate(groups.values(), 1):
        box = [min(w["x"] for w in words), min(w["y"] for w in words),
               max(w["x"] + w["w"] for w in words), max(w["y"] + w["h"] for w in words)]
        if not result.add(" ".join(w["text"] for w in words),
                          {"kind": "pdf", "page": page, "line": index,
                           "bbox": box, "coordinateSpace": "raster-pixels"},
                          method="ocr", confidence=min(w["confidence"] for w in words)):
            break
    image_path.unlink(missing_ok=True)
    if not groups:
        result.problem("PAGE_WITHOUT_TEXT")


def extract_pdf(source: Path, work: Path, budget: Budget, result: Result):
    info = run_tool("pdfinfo", [str(source)], work, budget).decode("utf-8", "replace")
    match = re.search(r"^Pages:\s+(\d+)\s*$", info, re.M)
    if not match or int(match[1]) < 1:
        raise ExtractionError("INVALID_PDF")
    pages = int(match[1])
    if pages > MAX_PAGES:
        result.problem("PAGE_LIMIT")
    image_pages = set()
    try:
        images = run_tool("pdfimages", ["-f", "1", "-l", str(min(pages, MAX_PAGES)),
                          "-list", str(source)], work, budget).decode("utf-8", "replace")
        for row in images.splitlines():
            columns = row.split()
            if len(columns) >= 3 and columns[0].isdigit() and columns[1].isdigit():
                image_pages.add(int(columns[0]))
    except ExtractionError:
        result.problem("PDF_IMAGE_COVERAGE_UNVERIFIED")
    for page in range(1, min(pages, MAX_PAGES) + 1):
        budget.remaining()
        if result.text_chars >= MAX_TEXT or len(result.units) >= MAX_UNITS:
            result.problem("TEXT_LIMIT" if result.text_chars >= MAX_TEXT else "UNIT_LIMIT")
            break
        try:
            data = run_tool("pdftotext", ["-f", str(page), "-l", str(page),
                            "-bbox-layout", "-enc", "UTF-8", str(source), "-"], work, budget)
            document = parse_xml(data, generated=True)
            lines = document.findall(".//{*}line")
            present = any("".join(word.itertext()).strip()
                          for word in document.findall(".//{*}word"))
            if not present:
                ocr_page(source, page, work, budget, result)
                continue
            if page in image_pages:
                # User requested OCR only for empty text layers. Mixed pages may
                # still contain text inside images, so do not claim full coverage.
                result.problem("PDF_IMAGES_NOT_OCR")
            if not lines:
                result.problem("PDF_LAYOUT_UNAVAILABLE")
                lines = document.findall(".//{*}page")
            for index, line in enumerate(lines, 1):
                text = " ".join("".join(word.itertext()) for word in line.findall(".//{*}word"))
                box = rectangle(line) if "xMin" in line.attrib else None
                if not result.add(text, {"kind": "pdf", "page": page, "line": index,
                                        "bbox": box, "coordinateSpace": "pdf-points"},
                                  method="native"):
                    break
        except ExtractionError as error:
            result.problem(str(error))
            # Other pages can still provide grounded positive facts.
            if str(error) in ("TIME_LIMIT", "DEPENDENCY_PDFTOTEXT_MISSING"):
                break


def paragraph_text(element) -> str:
    def pieces(node):
        if node is not element and node.tag in (W + "p", W + "del"):
            return
        if node.tag == W + "t":
            yield node.text or ""
        elif node.tag == W + "tab":
            yield "\t"
        elif node.tag in (W + "br", W + "cr"):
            yield "\n"
        for child in node:
            yield from pieces(child)
    return "".join(pieces(element))


def extract_docx(source: Path, budget: Budget, result: Result):
    with inspect_zip(source) as archive:
        names = set(archive.namelist())
        parts = ["word/document.xml"]
        document = parse_xml(archive.read(parts[0]))
        references = {element.attrib.get(REL + "id") for element in document.iter()
                      if element.tag in (W + "headerReference", W + "footerReference")}
        wanted_notes = {"footnote": {node.attrib.get(W + "id") for node in document.iter(W + "footnoteReference")},
                        "endnote": {node.attrib.get(W + "id") for node in document.iter(W + "endnoteReference")}}
        relationships = "word/_rels/document.xml.rels"
        if relationships in names:
            for rel in parse_xml(archive.read(relationships)):
                target = rel.attrib.get("Target", "")
                kind = rel.attrib.get("Type", "").rsplit("/", 1)[-1]
                required = rel.attrib.get("Id") in references or (
                    kind in ("footnotes", "endnotes") and wanted_notes[kind[:-1]])
                if required:
                    if (rel.attrib.get("TargetMode") == "External" or target.startswith("/")
                            or ".." in target.split("/") or "\\" in target):
                        result.problem("DOCX_UNRESOLVED_PART")
                    elif "word/" + target not in names:
                        result.problem("DOCX_UNRESOLVED_PART")
                    else:
                        parts.append("word/" + target)
                    references.discard(rel.attrib.get("Id"))
        if references:
            result.problem("DOCX_UNRESOLVED_PART")
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            result.problem("UNSUPPORTED_ACTIVE_CONTENT")
        for part in dict.fromkeys(parts):
            budget.remaining()
            tree = document if part == parts[0] else parse_xml(archive.read(part))
            local_tags = {node.tag.rsplit("}", 1)[-1] for node in tree.iter()}
            if local_tags.intersection({"ins", "del", "moveFrom", "moveTo"}):
                result.problem("DOCX_TRACKED_CHANGES")
            if local_tags.intersection({"drawing", "pict", "object", "altChunk", "oMath", "oMathPara"}):
                result.problem("DOCX_UNSUPPORTED_CONTENT")
            if local_tags.intersection({"fldSimple", "instrText"}):
                result.problem("DOCX_FIELDS_NOT_RECALCULATED")

            def walk(node, location):
                budget.remaining()
                tag = node.tag.rsplit("}", 1)[-1]
                if tag in wanted_notes and node.attrib.get(W + "id") not in wanted_notes[tag]:
                    return True
                if node.tag == W + "del":
                    return True
                if node.tag == W + "p" and not result.add(
                        paragraph_text(node), {"kind": "docx", "part": part, "path": location},
                        method="native"):
                    return False
                counts: dict[str, int] = {}
                for child in node:
                    child_tag = child.tag.rsplit("}", 1)[-1]
                    counts[child_tag] = counts.get(child_tag, 0) + 1
                    if not walk(child, location + "/" + child_tag + "[" + str(counts[child_tag]) + "]"):
                        return False
                return True
            if not walk(tree, "/" + tree.tag.rsplit("}", 1)[-1] + "[1]"):
                break
        for kind, identifiers in wanted_notes.items():
            if identifiers and not any(part.endswith(kind + "s.xml") for part in parts):
                result.problem("DOCX_UNRESOLVED_PART")


def cell_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ExtractionError("INVALID_CELL_VALUE")
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    raise ExtractionError("UNSUPPORTED_CELL_VALUE")


def extract_xlsx(source: Path, budget: Budget, result: Result):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ExtractionError("DEPENDENCY_OPENPYXL_MISSING") from None
    with inspect_zip(source) as archive:
        names = archive.namelist()
        if any(name.startswith(("xl/drawings/", "xl/embeddings/", "xl/charts/"))
               or name.lower().endswith("vbaproject.bin") for name in names):
            result.problem("XLSX_UNSUPPORTED_CONTENT")
        if any(name.startswith("xl/externalLinks/") for name in names):
            result.problem("XLSX_EXTERNAL_REFERENCES")
        # Reject DTDs before third-party XML parsing, including shared strings.
        for name in names:
            budget.remaining()
            if name.endswith((".xml", ".rels")):
                data = archive.read(name)
                if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                    raise ExtractionError("UNSAFE_XML")
                if name.startswith("xl/worksheets/") and b"headerFooter" in data:
                    tree = parse_xml(data)
                    if any("".join(node.itertext()).strip()
                           for node in tree.findall(".//{*}headerFooter")):
                        result.problem("XLSX_HEADERS_FOOTERS_NOT_EXTRACTED")
                    del tree
                del data
    with source.open("rb") as formula_stream, source.open("rb") as cached_stream:
        formulas = load_workbook(formula_stream, read_only=True, data_only=False,
                                 keep_links=False, keep_vba=False)
        cached = None
        try:
            cached = load_workbook(cached_stream, read_only=True, data_only=True,
                                   keep_links=False, keep_vba=False)
            if len(formulas.worksheets) > 100:
                result.problem("SHEET_LIMIT")
            visited = 0
            for worksheet in formulas.worksheets[:100]:
                values = cached[worksheet.title]
                # Do not trust producer-supplied dimensions (can hide real cells).
                worksheet.reset_dimensions()
                values.reset_dimensions()
                from itertools import zip_longest
                for row, cached_row in zip_longest(worksheet.rows, values.rows):
                    budget.remaining()
                    if row is None or cached_row is None or len(row) != len(cached_row):
                        raise ExtractionError("XLSX_PASS_MISMATCH")
                    for cell, cache in zip(row, cached_row):
                        visited += 1
                        if visited > MAX_CELLS:
                            result.problem("CELL_LIMIT")
                            return
                        if cell.value is None:
                            continue
                        formula = cell.value if cell.data_type == "f" else None
                        value = cell_value(cache.value if formula else cell.value)
                        if formula:
                            result.problem("XLSX_FORMULAS_NOT_RECALCULATED")
                            if value is None:
                                result.problem("XLSX_FORMULA_CACHE_MISSING")
                        if cell.data_type == "e" or cache.data_type == "e":
                            result.problem("XLSX_CELL_ERROR")
                        text = str(value) if value is not None else str(formula)
                        if not result.add(text, {"kind": "xlsx", "sheet": worksheet.title,
                                                "cell": cell.coordinate},
                                          method="native", value=value, formula=formula,
                                          numberFormat=cell.number_format,
                                          sheetState=worksheet.sheet_state):
                            return
        finally:
            formulas.close()
            if cached:
                cached.close()


def atomic_private(path: Path, data: bytes):
    if len(data) > MAX_OUTPUT or path.is_symlink():
        raise ExtractionError("OUTPUT_LIMIT" if len(data) > MAX_OUTPUT else "UNSAFE_PATH")
    descriptor, temporary = tempfile.mkstemp(prefix=".write-", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        if os.name == "posix":
            os.chmod(path, 0o600)
    finally:
        Path(temporary).unlink(missing_ok=True)


def cached_result(folder: Path, digest: str) -> tuple[dict, Path] | None:
    index = folder / (digest + ".cache.json")
    if index.is_symlink():
        raise ExtractionError("UNSAFE_PATH")
    try:
        with index.open("rb") as stream:
            reference = json.loads(bounded_read(stream, 1024))
        result_hash = reference["resultSha256"]
        if not re.fullmatch("[a-f0-9]{64}", result_hash):
            return None
        output = folder / (digest + "." + result_hash + ".json")
        if output.is_symlink():
            raise ExtractionError("UNSAFE_PATH")
        with output.open("rb") as stream:
            data = bounded_read(stream, MAX_OUTPUT)
        if hashlib.sha256(data).hexdigest() != result_hash:
            return None
        value = json.loads(data)
        if (value["sourceSha256"] == digest and value["extractorVersion"] == EXTRACTOR_VERSION
                and value["status"] == "COMPLETE" and not value["problems"]):
            return value, output
    except (OSError, ValueError, KeyError, TypeError):
        pass
    return None


def summary(value: dict, output: Path, cache_hit: bool) -> dict:
    return {"ok": True, "status": value["status"], "sha256": value["sourceSha256"],
            "extractorVersion": EXTRACTOR_VERSION, "mimeType": value["mimeType"],
            "outputPath": str(output), "unitCount": len(value["units"]),
            "textChars": value["textChars"], "problems": value["problems"],
            "cacheHit": cache_hit}


def extract(request: dict, budget: Budget | None = None) -> dict:
    budget = budget or Budget()
    root, source, digest, mime = checked_source(request)
    cache_root = root / ".extracted"
    private_directory(cache_root)
    folder = cache_root / EXTRACTOR_VERSION
    private_directory(folder)
    with tempfile.TemporaryDirectory(prefix=".extract-", dir=folder) as temporary:
        work = Path(temporary)
        copy = work / "source.bin"
        copy_verified(source, copy, digest, budget)
        kind = identify(copy)
        mime_mismatch = mime not in (None, "", "application/octet-stream", MIMES.get(kind))
        existing = None if mime_mismatch else cached_result(folder, digest)
        if existing:
            return summary(*existing, True)
        result = Result(digest, kind)
        if mime_mismatch:
            result.problem("MIME_MISMATCH")
        try:
            if kind == "pdf":
                extract_pdf(copy, work, budget, result)
            elif kind == "docx":
                extract_docx(copy, budget, result)
            elif kind == "xlsx":
                extract_xlsx(copy, budget, result)
            else:
                result.problem("UNSUPPORTED_FORMAT")
        except ExtractionError as error:
            result.problem(str(error))
        except MemoryError:
            result.problem("MEMORY_LIMIT")
        except Exception:
            result.problem("DOCUMENT_PARSE_FAILED")
        value = result.payload()
        data = canonical_json(value)
        result_hash = hashlib.sha256(data).hexdigest()
        output = folder / (digest + "." + result_hash + ".json")
        atomic_private(output, data)
        if value["status"] == "COMPLETE":
            atomic_private(folder / (digest + ".cache.json"),
                           canonical_json({"resultSha256": result_hash}))
        return summary(value, output, False)


def apply_limits():
    os.umask(0o077)
    if os.name != "posix":
        return
    import resource
    for key, limit in ((resource.RLIMIT_AS, MEMORY_BYTES),
                       (resource.RLIMIT_CPU, TIMEOUT_SECONDS),
                       (resource.RLIMIT_FSIZE, MAX_OUTPUT),
                       (resource.RLIMIT_NOFILE, 96)):
        soft, hard = resource.getrlimit(key)
        effective = min(limit, hard) if hard != resource.RLIM_INFINITY else limit
        resource.setrlimit(key, (effective, hard))
    # Covers stalled stdin / Python parsers as well as subprocesses.
    def deadline(_signum, _frame):
        raise ExtractionError("TIME_LIMIT")
    signal.signal(signal.SIGALRM, deadline)
    def cancelled(_signum, _frame):
        raise ExtractionError("PROCESS_CANCELLED")
    signal.signal(signal.SIGTERM, cancelled)
    signal.signal(signal.SIGINT, cancelled)
    signal.alarm(TIMEOUT_SECONDS)


def main() -> int:
    try:
        apply_limits()
        request = json.loads(bounded_read(sys.stdin.buffer, MAX_REQUEST))
        # Third-party parser warnings can contain customer-controlled names.
        # Only our fixed error codes and metadata may reach process output.
        with open(os.devnull, "w") as sink:
            with contextlib.redirect_stdout(sink), contextlib.redirect_stderr(sink):
                response = extract(request)
    except ExtractionError as error:
        response = {"ok": False, "status": "ERROR", "errorCode": str(error),
                    "extractorVersion": EXTRACTOR_VERSION}
    except MemoryError:
        response = {"ok": False, "status": "ERROR", "errorCode": "MEMORY_LIMIT",
                    "extractorVersion": EXTRACTOR_VERSION}
    except Exception:
        response = {"ok": False, "status": "ERROR", "errorCode": "INVALID_OR_UNAVAILABLE_SOURCE",
                    "extractorVersion": EXTRACTOR_VERSION}
    if os.name == "posix":
        signal.alarm(0)
    sys.stdout.write(json.dumps(response, ensure_ascii=True, allow_nan=False) + "\n")
    return 0 if response["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
